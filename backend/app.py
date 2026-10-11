"""
CodeArena — Python Flask Backend
Handles: users, problems, contests, submissions, leaderboard
Database: Excel workbook (data/contest.xlsx) via openpyxl
Queue   : Redis (pushes to judge-worker) with built-in standalone local judge
"""

import json
import os
import uuid
import hashlib
import secrets
import time
import socket
import threading
import subprocess
import tempfile
import sys
import shutil

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')

from datetime import datetime, timedelta
from functools import wraps

import redis
from flask import Flask, request, jsonify, g, send_file
from flask_cors import CORS
import openpyxl
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment

app = Flask(__name__)
CORS(app, origins="*", supports_credentials=True)

REDIS_ADDR   = os.getenv("REDIS_ADDR", "127.0.0.1:6379")
REDIS_QUEUE  = os.getenv("REDIS_QUEUE", "submissions")
REDIS_RESULT = os.getenv("REDIS_RESULT_PREFIX", "result:")
DATA_FILE    = os.getenv("DATA_FILE", "data/contest.xlsx")
SECRET_KEY   = os.getenv("SECRET_KEY", "codearena-secret-2026")

WB_LOCK = threading.Lock()
LOCAL_RESULTS = {}

redis_host, redis_port = REDIS_ADDR.rsplit(":", 1)
REDIS_OK = False
rdb = None
try:
    with socket.create_connection((redis_host, int(redis_port)), timeout=0.2):
        rdb = redis.Redis(host=redis_host, port=int(redis_port), decode_responses=True, socket_connect_timeout=1, socket_timeout=1)
        rdb.ping()
        REDIS_OK = True
except Exception:
    REDIS_OK = False
    rdb = None

SHEETS = ["users","problems","contests","contest_problems","submissions","leaderboard","sessions"]
HEADERS = {
    "users":            ["id","username","email","password_hash","role","rating","created_at","avatar_color"],
    "problems":         ["id","title","slug","difficulty","time_limit_ms","mem_limit_mb","statement","input_format","output_format","constraints","sample_input","sample_output","tags","test_cases_json","author_id","created_at"],
    "contests":         ["id","title","slug","start_time","end_time","status","description","is_public","created_by","created_at"],
    "contest_problems": ["id","contest_id","problem_id","order","points"],
    "submissions":      ["id","user_id","problem_id","contest_id","language","code","verdict","score","time_ms","memory_mb","submitted_at","judge_result_json"],
    "leaderboard":      ["id","contest_id","user_id","username","total_score","problems_solved","last_accepted_at","rank"],
    "sessions":         ["token","user_id","created_at","expires_at"],
}

def _now():
    return datetime.utcnow().isoformat()

def _hash_pw(pw):
    return hashlib.sha256((pw + SECRET_KEY).encode()).hexdigest()

def _make_token():
    return secrets.token_urlsafe(32)

def ensure_workbook():
    os.makedirs("data", exist_ok=True)
    with WB_LOCK:
        if not os.path.exists(DATA_FILE):
            wb = Workbook()
            wb.remove(wb.active)
            for sheet, headers in HEADERS.items():
                ws = wb.create_sheet(sheet)
                ws.append(headers)
                for cell in ws[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="7C3AED")
                    cell.alignment = Alignment(horizontal="center")
            wb.save(DATA_FILE)
            _seed_demo_data(wb)

def load_wb():
    with WB_LOCK:
        return load_workbook(DATA_FILE)

def save_wb(wb):
    with WB_LOCK:
        wb.save(DATA_FILE)

def sheet_to_dicts(wb, sheet_name):
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = rows[0]
    return [dict(zip(headers, row)) for row in rows[1:] if any(v is not None for v in row)]

def append_row(wb, sheet_name, row_dict):
    ws = wb[sheet_name]
    headers = HEADERS[sheet_name]
    ws.append([row_dict.get(h) for h in headers])

def update_row(wb, sheet_name, match_key, match_val, updates):
    ws = wb[sheet_name]
    headers = HEADERS[sheet_name]
    key_col = headers.index(match_key) + 1
    for row in ws.iter_rows(min_row=2):
        if str(row[key_col-1].value) == str(match_val):
            for k, v in updates.items():
                if k in headers:
                    row[headers.index(k)].value = v
            return True
    return False

def delete_row(wb, sheet_name, match_key, match_val):
    ws = wb[sheet_name]
    headers = HEADERS[sheet_name]
    key_col = headers.index(match_key) + 1
    for row in ws.iter_rows(min_row=2):
        if str(row[key_col-1].value) == str(match_val):
            ws.delete_rows(row[0].row)
            return True
    return False

def _seed_demo_data(wb=None):  # noqa: C901
    should_save = False
    if wb is None:
        wb = load_wb()
        should_save = True

    uid = str(uuid.uuid4())
    append_row(wb, "users", {"id":uid,"username":"admin","email":"admin@codearena.io","password_hash":_hash_pw("admin123"),"role":"admin","rating":2500,"created_at":_now(),"avatar_color":"#7c3aed"})
    
    problems_data = [
        {'title': 'Two Sum', 'slug': 'two-sum', 'difficulty': 'Easy', 'time_limit_ms': 1000, 'mem_limit_mb': 256, 'statement': 'Given an array of integers nums and an integer target, return indices of the two numbers such that they add up to target.', 'input_format': 'First line: n and target. Second line: n space-separated integers.', 'output_format': 'Two space-separated 0-based indices.', 'constraints': '2 <= n <= 10^4, -10^9 <= nums[i] <= 10^9', 'sample_input': '4 9\n2 7 11 15', 'sample_output': '0 1', 'tags': 'array,hash-map', 'test_cases': [{'index': 0, 'input': '4 9\n2 7 11 15', 'expected': '0 1'}, {'index': 1, 'input': '3 6\n3 2 4', 'expected': '1 2'}]},
        {'title': 'Fibonacci Sequence', 'slug': 'fibonacci', 'difficulty': 'Easy', 'time_limit_ms': 1000, 'mem_limit_mb': 256, 'statement': 'Print the first n Fibonacci numbers starting from 0.', 'input_format': 'Single integer n.', 'output_format': 'Space-separated Fibonacci numbers.', 'constraints': '1 <= n <= 50', 'sample_input': '5', 'sample_output': '0 1 1 2 3', 'tags': 'dp,math', 'test_cases': [{'index': 0, 'input': '5', 'expected': '0 1 1 2 3'}, {'index': 1, 'input': '1', 'expected': '0'}]},
        {'title': 'Longest Common Subsequence', 'slug': 'lcs', 'difficulty': 'Medium', 'time_limit_ms': 2000, 'mem_limit_mb': 256, 'statement': 'Find the length of the longest common subsequence between two strings.', 'input_format': 'Two lines, each with a string.', 'output_format': 'Single integer LCS length.', 'constraints': '1 <= |s|, |t| <= 1000', 'sample_input': 'abcde\nace', 'sample_output': '3', 'tags': 'dp,string', 'test_cases': [{'index': 0, 'input': 'abcde\nace', 'expected': '3'}, {'index': 1, 'input': 'abc\nabc', 'expected': '3'}]},
        {'title': 'Maximum Subarray', 'slug': 'max-subarray', 'difficulty': 'Medium', 'time_limit_ms': 1000, 'mem_limit_mb': 256, 'statement': 'Find the contiguous subarray with the largest sum (Kadanes Algorithm).', 'input_format': 'First line: n. Second line: n integers.', 'output_format': 'Maximum subarray sum.', 'constraints': '1 <= n <= 10^5', 'sample_input': '9\n-2 1 -3 4 -1 2 1 -5 4', 'sample_output': '6', 'tags': 'dp,array', 'test_cases': [{'index': 0, 'input': '9\n-2 1 -3 4 -1 2 1 -5 4', 'expected': '6'}, {'index': 1, 'input': '1\n1', 'expected': '1'}]},
        {'title': 'Binary Search', 'slug': 'binary-search', 'difficulty': 'Easy', 'time_limit_ms': 1000, 'mem_limit_mb': 256, 'statement': 'Find index of target in sorted array. Return -1 if not found.', 'input_format': 'Line 1: n and target. Line 2: n sorted integers.', 'output_format': 'Index of target or -1.', 'constraints': '1 <= n <= 10^4', 'sample_input': '6 9\n1 3 5 7 9 11', 'sample_output': '4', 'tags': 'binary-search,array', 'test_cases': [{'index': 0, 'input': '6 9\n1 3 5 7 9 11', 'expected': '4'}, {'index': 1, 'input': '3 0\n1 2 3', 'expected': '-1'}]}
    ]

    pids = []
    for p in problems_data:
        pid = str(uuid.uuid4())
        pids.append(pid)
        append_row(wb, "problems", {"id":pid,"title":p["title"],"slug":p["slug"],"difficulty":p["difficulty"],"time_limit_ms":p["time_limit_ms"],"mem_limit_mb":p["mem_limit_mb"],"statement":p["statement"],"input_format":p["input_format"],"output_format":p["output_format"],"constraints":p["constraints"],"sample_input":p["sample_input"],"sample_output":p["sample_output"],"tags":p["tags"],"test_cases_json":json.dumps(p["test_cases"]),"author_id":uid,"created_at":_now()})
    cid = str(uuid.uuid4())
    now = datetime.utcnow()
    append_row(wb, "contests", {"id":cid,"title":"CodeArena Round #1","slug":"round-1","start_time":now.isoformat(),"end_time":(now+timedelta(hours=3)).isoformat(),"status":"ongoing","description":"Welcome to CodeArena! Compete, code, and climb the leaderboard.","is_public":True,"created_by":uid,"created_at":_now()})
    points = [100,100,200,200,300]
    for i, pid in enumerate(pids):
        append_row(wb, "contest_problems", {"id":str(uuid.uuid4()),"contest_id":cid,"problem_id":pid,"order":i+1,"points":points[i]})
    if should_save:
        save_wb(wb)

def _normalize_output(text):
    if not text:
        return ""
    return "\n".join([line.rstrip() for line in str(text).strip().splitlines()])

def _execute_test_case(language, code, input_data, time_limit_ms, temp_dir):
    timeout_sec = max(0.8, (time_limit_ms or 1000) / 1000.0)
    lang = (language or "").lower().strip()
    start_t = time.time()
    try:
        if lang in ("python", "py", "python3"):
            src_file = os.path.join(temp_dir, "solution.py")
            with open(src_file, "w", encoding="utf-8") as f:
                f.write(code)
            proc = subprocess.run([sys.executable, src_file], input=input_data, capture_output=True, text=True, timeout=timeout_sec)
            elapsed_ms = int((time.time() - start_t) * 1000)
            if proc.returncode != 0:
                return {"verdict": "RE", "time_ms": elapsed_ms, "error": proc.stderr[:300]}
            return {"verdict": "OK", "time_ms": elapsed_ms, "output": proc.stdout}

        elif lang in ("javascript", "js", "node"):
            src_file = os.path.join(temp_dir, "solution.js")
            with open(src_file, "w", encoding="utf-8") as f:
                f.write(code)
            proc = subprocess.run(["node", src_file], input=input_data, capture_output=True, text=True, timeout=timeout_sec)
            elapsed_ms = int((time.time() - start_t) * 1000)
            if proc.returncode != 0:
                return {"verdict": "RE", "time_ms": elapsed_ms, "error": proc.stderr[:300]}
            return {"verdict": "OK", "time_ms": elapsed_ms, "output": proc.stdout}

        elif lang in ("cpp", "c++"):
            src_file = os.path.join(temp_dir, "solution.cpp")
            exe_file = os.path.join(temp_dir, "solution.exe" if os.name == "nt" else "solution")
            with open(src_file, "w", encoding="utf-8") as f:
                f.write(code)
            comp = subprocess.run(["g++", "-O2", src_file, "-o", exe_file], capture_output=True, text=True, timeout=10.0)
            if comp.returncode != 0:
                return {"verdict": "CE", "time_ms": 0, "error": comp.stderr[:300]}
            proc = subprocess.run([exe_file], input=input_data, capture_output=True, text=True, timeout=timeout_sec)
            elapsed_ms = int((time.time() - start_t) * 1000)
            if proc.returncode != 0:
                return {"verdict": "RE", "time_ms": elapsed_ms, "error": proc.stderr[:300]}
            return {"verdict": "OK", "time_ms": elapsed_ms, "output": proc.stdout}

        else:
            return {"verdict": "CE", "time_ms": 0, "error": f"Language {language} execution not configured"}

    except subprocess.TimeoutExpired:
        return {"verdict": "TLE", "time_ms": int(timeout_sec * 1000), "error": "Time Limit Exceeded"}
    except Exception as e:
        return {"verdict": "RE", "time_ms": 0, "error": str(e)[:300]}

def _update_leaderboard(wb, contest_id, user_id, problem_id):
    users = sheet_to_dicts(wb, "users")
    user = next((u for u in users if str(u["id"]) == str(user_id)), None)
    username = user["username"] if user else "Coder"
    
    cps = sheet_to_dicts(wb, "contest_problems")
    cp = next((x for x in cps if str(x.get("contest_id")) == str(contest_id) and str(x.get("problem_id")) == str(problem_id)), None)
    pts = int(cp.get("points") or 100) if cp else 100

    lb = sheet_to_dicts(wb, "leaderboard")
    entry = next((e for e in lb if str(e.get("contest_id")) == str(contest_id) and str(e.get("user_id")) == str(user_id)), None)
    now_iso = _now()
    if entry:
        new_score = int(entry.get("total_score") or 0) + pts
        new_solved = int(entry.get("problems_solved") or 0) + 1
        update_row(wb, "leaderboard", "id", entry["id"], {
            "total_score": new_score,
            "problems_solved": new_solved,
            "last_accepted_at": now_iso
        })
    else:
        append_row(wb, "leaderboard", {
            "id": str(uuid.uuid4()),
            "contest_id": contest_id,
            "user_id": user_id,
            "username": username,
            "total_score": pts,
            "problems_solved": 1,
            "last_accepted_at": now_iso,
            "rank": 1
        })

def _judge_submission_local(sid, user_id, contest_id, problem_id, language, code, test_cases, time_limit_ms, mem_limit_mb):
    temp_dir = tempfile.mkdtemp(prefix=f"judge_{sid[:8]}_")
    try:
        final_verdict = "AC"
        max_time = 0
        test_results = []
        if not test_cases:
            test_cases = [{"index": 0, "input": "", "expected": ""}]

        for idx, tc in enumerate(test_cases):
            inp = str(tc.get("input", ""))
            exp = str(tc.get("expected", ""))
            res = _execute_test_case(language, code, inp, time_limit_ms, temp_dir)
            t_ms = res.get("time_ms", 0)
            max_time = max(max_time, t_ms)

            if res["verdict"] == "CE":
                final_verdict = "CE"
                test_results.append({"test_index": idx, "verdict": "CE", "error": res.get("error")})
                break
            elif res["verdict"] == "TLE":
                final_verdict = "TLE"
                test_results.append({"test_index": idx, "verdict": "TLE", "time_ms": t_ms})
                break
            elif res["verdict"] == "RE":
                final_verdict = "RE"
                test_results.append({"test_index": idx, "verdict": "RE", "time_ms": t_ms, "error": res.get("error")})
                break
            elif res["verdict"] == "OK":
                actual = _normalize_output(res.get("output", ""))
                expected = _normalize_output(exp)
                if exp and actual != expected:
                    final_verdict = "WA"
                    test_results.append({"test_index": idx, "verdict": "WA", "time_ms": t_ms, "actual": actual[:100], "expected": expected[:100]})
                    break
                else:
                    test_results.append({"test_index": idx, "verdict": "AC", "time_ms": t_ms})

        result_payload = {
            "submission_id": sid,
            "verdict": final_verdict,
            "max_time_ms": max_time,
            "max_memory_mb": 16,
            "tests_passed": len([t for t in test_results if t.get("verdict") == "AC"]),
            "tests_total": len(test_cases),
            "results": test_results
        }
        LOCAL_RESULTS[sid] = result_payload

        wb = load_wb()
        update_row(wb, "submissions", "id", sid, {
            "verdict": final_verdict,
            "score": 100 if final_verdict == "AC" else 0,
            "time_ms": max_time,
            "memory_mb": 16,
            "judge_result_json": json.dumps(result_payload)
        })
        if final_verdict == "AC" and contest_id:
            _update_leaderboard(wb, contest_id, user_id, problem_id)
        save_wb(wb)
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)

def _get_user_from_token(token):
    if not token:
        return None
    wb = load_wb()
    sessions = sheet_to_dicts(wb, "sessions")
    for s in sessions:
        if s["token"] == token:
            exp = s.get("expires_at")
            if exp and datetime.fromisoformat(str(exp)) > datetime.utcnow():
                users = sheet_to_dicts(wb, "users")
                return next((u for u in users if str(u["id"]) == str(s["user_id"])), None)
    return None

def require_auth(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        token = request.headers.get("Authorization","").replace("Bearer ","")
        user = _get_user_from_token(token)
        if not user:
            return jsonify({"error":"Unauthorized"}), 401
        g.current_user = user
        return f(*args, **kwargs)
    return decorated

def require_admin(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        token = request.headers.get("Authorization","").replace("Bearer ","")
        user = _get_user_from_token(token)
        if not user or user.get("role") != "admin":
            return jsonify({"error":"Forbidden"}), 403
        g.current_user = user
        return f(*args, **kwargs)
    return decorated

@app.route("/api/auth/register", methods=["POST"])
def register():
    data = request.json or {}
    username = str(data.get("username","")).strip()
    email    = str(data.get("email","")).strip()
    password = str(data.get("password",""))
    if not username or not email or not password:
        return jsonify({"error":"All fields required"}), 400
    wb = load_wb()
    users = sheet_to_dicts(wb, "users")
    if any(u["username"] == username for u in users):
        return jsonify({"error":"Username already taken"}), 409
    if any(u["email"] == email for u in users):
        return jsonify({"error":"Email already registered"}), 409
    colors = ["#7c3aed","#06b6d4","#f59e0b","#10b981","#ef4444","#8b5cf6"]
    uid = str(uuid.uuid4())
    append_row(wb,"users",{"id":uid,"username":username,"email":email,"password_hash":_hash_pw(password),"role":"user","rating":1200,"created_at":_now(),"avatar_color":colors[len(users)%len(colors)]})
    save_wb(wb)
    return jsonify({"message":"Registered"}), 201

@app.route("/api/auth/login", methods=["POST"])
def login():
    data = request.json or {}
    username = str(data.get("username","")).strip()
    password = str(data.get("password",""))
    wb = load_wb()
    users = sheet_to_dicts(wb, "users")
    user = next((u for u in users if u["username"] == username and u["password_hash"] == _hash_pw(password)), None)
    if not user:
        return jsonify({"error":"Invalid credentials"}), 401
    token = _make_token()
    expires = (datetime.utcnow()+timedelta(days=7)).isoformat()
    sessions = sheet_to_dicts(wb, "sessions")
    ws = wb["sessions"]
    rows_to_del = [i+2 for i,s in enumerate(sessions) if str(s["user_id"]) == str(user["id"])]
    for r in reversed(rows_to_del):
        ws.delete_rows(r)
    append_row(wb,"sessions",{"token":token,"user_id":user["id"],"created_at":_now(),"expires_at":expires})
    save_wb(wb)
    return jsonify({"token":token,"user":{k:user[k] for k in ["id","username","email","role","rating","avatar_color"]}})

@app.route("/api/auth/logout", methods=["POST"])
@require_auth
def logout():
    token = request.headers.get("Authorization","").replace("Bearer ","")
    wb = load_wb()
    delete_row(wb, "sessions", "token", token)
    save_wb(wb)
    return jsonify({"message":"Logged out"})

@app.route("/api/auth/me", methods=["GET"])
@require_auth
def me():
    u = g.current_user
    return jsonify({"user":{k:u[k] for k in ["id","username","email","role","rating","avatar_color"]}})

@app.route("/api/problems", methods=["GET"])
def list_problems():
    wb = load_wb()
    problems = sheet_to_dicts(wb, "problems")
    result = []
    for p in problems:
        tags = (p.get("tags") or "").split(",") if p.get("tags") else []
        result.append({
            "id": p["id"], "title": p["title"], "slug": p["slug"],
            "difficulty": p["difficulty"], "tags": tags,
            "time_limit_ms": p.get("time_limit_ms", 1000),
            "mem_limit_mb": p.get("mem_limit_mb", 256)
        })
    return jsonify(result)

@app.route("/api/problems/<slug>", methods=["GET"])
def get_problem(slug):
    wb = load_wb()
    problems = sheet_to_dicts(wb, "problems")
    p = next((x for x in problems if x["slug"] == slug), None)
    if not p:
        return jsonify({"error":"Not found"}), 404
    result = {k:p[k] for k in p if k != "test_cases_json"}
    result["tags"] = (result.get("tags") or "").split(",")
    return jsonify(result)

@app.route("/api/problems", methods=["POST"])
@require_admin
def create_problem():
    data = request.json or {}
    wb = load_wb()
    pid = str(uuid.uuid4())
    append_row(wb,"problems",{"id":pid,"title":data["title"],"slug":data["slug"],"difficulty":data.get("difficulty","Easy"),"time_limit_ms":data.get("time_limit_ms",1000),"mem_limit_mb":data.get("mem_limit_mb",256),"statement":data.get("statement",""),"input_format":data.get("input_format",""),"output_format":data.get("output_format",""),"constraints":data.get("constraints",""),"sample_input":data.get("sample_input",""),"sample_output":data.get("sample_output",""),"tags":",".join(data.get("tags",[])),"test_cases_json":json.dumps(data.get("test_cases",[])),"author_id":g.current_user["id"],"created_at":_now()})
    save_wb(wb)
    return jsonify({"id":pid}), 201

@app.route("/api/contests", methods=["GET"])
def list_contests():
    wb = load_wb()
    contests = sheet_to_dicts(wb, "contests")
    return jsonify([{"id":c["id"],"title":c["title"],"slug":c["slug"],"start_time":c["start_time"],"end_time":c["end_time"],"status":c["status"],"description":c["description"]} for c in contests])

@app.route("/api/contests/<slug>", methods=["GET"])
def get_contest(slug):
    wb = load_wb()
    contests = sheet_to_dicts(wb, "contests")
    c = next((x for x in contests if x["slug"] == slug), None)
    if not c:
        return jsonify({"error":"Not found"}), 404
    cp = sheet_to_dicts(wb,"contest_problems")
    problems = sheet_to_dicts(wb,"problems")
    cps = sorted([r for r in cp if str(r["contest_id"]) == str(c["id"])], key=lambda x: int(x.get("order") or 0))
    prob_list = []
    for cp_row in cps:
        p = next((px for px in problems if str(px["id"]) == str(cp_row["problem_id"])), None)
        if p:
            prob_list.append({"id":p["id"],"title":p["title"],"slug":p["slug"],"difficulty":p["difficulty"],"points":cp_row["points"],"order":cp_row["order"]})
    return jsonify({**{k:c[k] for k in ["id","title","slug","start_time","end_time","status","description"]},"problems":prob_list})

@app.route("/api/contests", methods=["POST"])
@require_admin
def create_contest():
    data = request.json or {}
    cid = str(uuid.uuid4())
    wb = load_wb()
    append_row(wb,"contests",{"id":cid,"title":data["title"],"slug":data["slug"],"start_time":data["start_time"],"end_time":data["end_time"],"status":data.get("status","upcoming"),"description":data.get("description",""),"is_public":data.get("is_public",True),"created_by":g.current_user["id"],"created_at":_now()})
    save_wb(wb)
    return jsonify({"id":cid}), 201

@app.route("/api/submissions", methods=["POST"])
@require_auth
def submit():
    data = request.json or {}
    problem_slug = data.get("problem_slug","")
    language = data.get("language","")
    code = data.get("code","")
    if not problem_slug or not language or not code:
        return jsonify({"error":"problem_slug, language, code required"}), 400
    wb = load_wb()
    problems = sheet_to_dicts(wb,"problems")
    problem = next((p for p in problems if p["slug"] == problem_slug), None)
    if not problem:
        return jsonify({"error":"Problem not found"}), 404
    sid = str(uuid.uuid4())
    test_cases = json.loads(problem.get("test_cases_json") or "[]")
    time_limit = int(problem.get("time_limit_ms") or 1000)
    mem_limit = int(problem.get("mem_limit_mb") or 256)
    verdict = "PENDING"

    append_row(wb,"submissions",{"id":sid,"user_id":g.current_user["id"],"problem_id":problem["id"],"contest_id":data.get("contest_id"),"language":language,"code":code,"verdict":verdict,"score":0,"time_ms":None,"memory_mb":None,"submitted_at":_now(),"judge_result_json":None})
    save_wb(wb)

    if REDIS_OK:
        payload = {"id":sid,"language":language,"code":code,"test_cases":test_cases,"time_limit_ms":time_limit,"mem_limit_mb":mem_limit}
        try:
            rdb.rpush(REDIS_QUEUE, json.dumps(payload))
        except Exception:
            th = threading.Thread(target=_judge_submission_local, args=(sid, g.current_user["id"], data.get("contest_id"), problem["id"], language, code, test_cases, time_limit, mem_limit))
            th.daemon = True
            th.start()
    else:
        th = threading.Thread(target=_judge_submission_local, args=(sid, g.current_user["id"], data.get("contest_id"), problem["id"], language, code, test_cases, time_limit, mem_limit))
        th.daemon = True
        th.start()

    return jsonify({"submission_id":sid,"verdict":verdict}), 202

@app.route("/api/submissions/<sid>", methods=["GET"])
@require_auth
def get_submission(sid):
    if sid in LOCAL_RESULTS:
        return jsonify(LOCAL_RESULTS[sid])

    if REDIS_OK:
        try:
            raw = rdb.get(REDIS_RESULT + sid)
            if raw:
                result = json.loads(raw)
                wb = load_wb()
                update_row(wb,"submissions","id",sid,{"verdict":result.get("verdict","SE"),"time_ms":result.get("max_time_ms",0),"memory_mb":result.get("max_memory_mb",0),"judge_result_json":raw})
                save_wb(wb)
                return jsonify(result)
        except Exception:
            pass

    wb = load_wb()
    subs = sheet_to_dicts(wb,"submissions")
    sub = next((s for s in subs if str(s["id"]) == sid), None)
    if not sub:
        return jsonify({"error":"Not found"}), 404

    raw_res = sub.get("judge_result_json")
    if raw_res:
        try:
            return jsonify(json.loads(raw_res))
        except Exception:
            pass

    return jsonify({"submission_id":sub["id"],"verdict":sub["verdict"],"time_ms":sub.get("time_ms") or 0,"memory_mb":sub.get("memory_mb") or 0,"language":sub["language"],"submitted_at":sub["submitted_at"]})

@app.route("/api/submissions", methods=["GET"])
def list_submissions():
    wb = load_wb()
    subs = sheet_to_dicts(wb,"submissions")
    problem_slug = request.args.get("problem_slug")
    user_id = request.args.get("user_id")
    limit = int(request.args.get("limit",30))
    if problem_slug:
        problems = sheet_to_dicts(wb,"problems")
        p = next((px for px in problems if px["slug"] == problem_slug), None)
        if p:
            subs = [s for s in subs if str(s["problem_id"]) == str(p["id"])]
    if user_id:
        subs = [s for s in subs if str(s["user_id"]) == str(user_id)]
    subs = sorted(subs, key=lambda x: str(x.get("submitted_at") or ""), reverse=True)[:limit]
    users = sheet_to_dicts(wb,"users")
    problems = sheet_to_dicts(wb,"problems")
    result = []
    for s in subs:
        u = next((x for x in users if str(x["id"]) == str(s["user_id"])), {})
        p = next((x for x in problems if str(x["id"] or "") == str(s.get("problem_id") or "")), {})
        result.append({"id":s["id"],"verdict":s["verdict"],"language":s["language"],"time_ms":s["time_ms"],"memory_mb":s["memory_mb"],"submitted_at":s["submitted_at"],"username":u.get("username","?"),"problem_title":p.get("title","?"),"problem_slug":p.get("slug","?")})
    return jsonify(result)

@app.route("/api/leaderboard", methods=["GET"])
def leaderboard():
    contest_slug = request.args.get("contest_slug")
    wb = load_wb()
    if contest_slug:
        contests = sheet_to_dicts(wb,"contests")
        c = next((x for x in contests if x["slug"] == contest_slug), None)
        if not c:
            return jsonify([])
        lb = sheet_to_dicts(wb,"leaderboard")
        rows = sorted([r for r in lb if str(r["contest_id"]) == str(c["id"])], key=lambda x: (-int(x.get("total_score") or 0), str(x.get("last_accepted_at") or "")))
        return jsonify(rows)
    users = sheet_to_dicts(wb,"users")
    subs = sheet_to_dicts(wb,"submissions")
    result = []
    for i, u in enumerate(sorted(users, key=lambda x: -(int(x.get("rating") or 0)))):
        solved = len(set(s["problem_id"] for s in subs if str(s["user_id"]) == str(u["id"]) and s["verdict"] == "AC"))
        result.append({"rank":i+1,"username":u["username"],"rating":u.get("rating",1200),"problems_solved":solved,"avatar_color":u.get("avatar_color","#7c3aed")})
    return jsonify(result)

@app.route("/api/stats", methods=["GET"])
def stats():
    wb = load_wb()
    return jsonify({"users":len(sheet_to_dicts(wb,"users")),"problems":len(sheet_to_dicts(wb,"problems")),"contests":len(sheet_to_dicts(wb,"contests")),"submissions":len(sheet_to_dicts(wb,"submissions"))})

@app.route("/api/health", methods=["GET"])
def health():
    return jsonify({"status":"ok","time":_now(),"redis":REDIS_OK,"database":"openpyxl Excel (contest.xlsx)"})

@app.route("/api/admin/export", methods=["GET"])
@require_admin
def export_excel():
    return send_file(os.path.abspath(DATA_FILE), as_attachment=True, download_name="contest_data.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

if __name__ == "__main__":
    ensure_workbook()
    print("\n  [CodeArena] Flask Backend starting on http://0.0.0.0:5000")
    print(f"  [CodeArena] Excel Database: {DATA_FILE}")
    print(f"  [CodeArena] Redis Status: {'Connected' if REDIS_OK else 'Offline (Standalone Local Judge Active)'}\n")
    app.run(host="0.0.0.0", port=5000, debug=False)
